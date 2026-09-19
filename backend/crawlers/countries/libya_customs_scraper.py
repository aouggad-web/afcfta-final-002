#!/usr/bin/env python3
"""
Crawler pour le tarif douanier Libye (XLSX officiel customs.gov.ly).
Source : https://customs.gov.ly/wp-content/uploads/2023/12/التعريفة-الجمركية-.2022-.xlsx
Format : XLSX avec colonnes en arabe :
  - رقم البند (numéro de position)
  - رمز النظام المنسق (code SH)
  - بيان المنتجات (désignation produit)
  - فئة الضريبة (taux du prélèvement du tarif)
  - التعريفة التفضيلية لدول جامعة الدول العربية (colonne préférentielle Ligue des États arabes)
"""

import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import httpx
import openpyxl

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
logger = logging.getLogger(__name__)

XLSX_URL = "https://customs.gov.ly/wp-content/uploads/2023/12/التعريفة-الجمركية-.2022-.xlsx"
DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "crawled"
TEMP_DIR = Path("/tmp/lby_tariff")

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36",
    "Accept": "*/*",
    "Accept-Language": "ar,en-US;q=0.9",
}


def download_xlsx() -> Optional[str]:
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    filepath = TEMP_DIR / "LBY_tariff.xlsx"
    if filepath.exists() and filepath.stat().st_size > 100000:
        logger.info(f"Using cached XLSX: {filepath}")
        return str(filepath)
    logger.info("Downloading Libya tariff XLSX...")
    resp = httpx.get(XLSX_URL, timeout=120, follow_redirects=True, verify=False, headers=HEADERS)
    if resp.status_code == 200 and len(resp.content) > 100000:
        with open(filepath, "wb") as f:
            f.write(resp.content)
        logger.info(f"Downloaded: {len(resp.content) // 1024} KB")
        return str(filepath)
    logger.error(f"Download failed: {resp.status_code}")
    return None


#: « ممنوع استيراده » — importation interdite. Ce n'est PAS un taux absent :
#: la marchandise ne peut pas entrer. La confondre avec une donnée manquante
#: ferait lire « droit indisponible » là où la réponse est « interdit », et
#: la traiter comme 0 % ferait liquider une importation prohibée.
MENTION_INTERDIT = "ممنوع استيراده"

#: « معفاة » — exonéré. Un zéro réellement publié par le tarif.
MENTION_EXONERE = "معفاة"


def parse_rate(val) -> Optional[float]:
    """Rend un taux en POURCENTS, ou ``None`` si la cellule n'en porte pas.

    Le tarif libyen mélange deux écritures dans la même colonne, relevé sur
    le XLSX officiel 2022 : une fraction décimale (``0.05``, 4 379 lignes) et
    un pourcentage littéral (``5%``, 356 lignes) — les deux valant 5 %. Une
    première version rendait ``float(s)`` pour les deux, soit ``0.05`` pris
    pour 0,05 % (cent fois trop bas) et une exception sur ``5%`` (taux
    perdu). Sa docstring annonçait pourtant la bonne conversion.

    La règle de lecture est donnée par les valeurs elles-mêmes : le tarif ne
    porte aucun droit inférieur à 1 %, et ``0.05``/``0.1``/``0.3`` se lisent
    5 %, 10 % et 30 % — ce que confirme la coexistence de ``5%``, ``10%`` et
    ``30%`` pour les mêmes marchandises. Une valeur ``≤ 1`` sans signe pour
    cent est donc une fraction ; au-delà, un pourcentage déjà écrit comme tel.

    ``None`` est rendu pour une cellule vide, pour une mention d'interdiction
    (traitée séparément par :func:`lire_cellule_droit`) et pour tout texte non
    reconnu — jamais un zéro de substitution.
    """
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    if MENTION_EXONERE in s or "exempt" in s.lower() or "free" in s.lower():
        return 0.0
    if MENTION_INTERDIT in s:
        return None
    pourcent = s.endswith("%")
    nombre = s.rstrip("%").strip()
    try:
        valeur = float(nombre)
    except ValueError:
        return None
    if pourcent:
        return valeur
    # Fraction décimale (0.05 → 5 %). Une valeur > 1 sans signe pour cent est
    # déjà un pourcentage : on la rend telle quelle plutôt que de la
    # multiplier par cent, ce qui produirait des droits de plusieurs milliers.
    return valeur * 100 if valeur <= 1 else valeur


def lire_cellule_droit(val):
    """Rend ``(taux_pct, interdit)`` pour une cellule de la colonne des droits.

    Sépare les trois réponses possibles du tarif — un taux, une interdiction
    d'importation, ou rien — pour qu'aucune ne soit lue pour une autre.
    """
    interdit = val is not None and MENTION_INTERDIT in str(val)
    return (None if interdit else parse_rate(val)), interdit


#: La feuille qui porte le tarif. Le classeur en compte trois, et la deuxième
#: est un piège : « ورقة2 » aligne 1 040 896 lignes dont 422 SEULEMENT portent
#: un code — un collecteur qui la lirait ramasserait 7 % du tarif en croyant
#: tout avoir. « ورقة3 » est vide. Prendre la feuille par son RANG marcherait
#: aujourd'hui et casserait en silence le jour où la douane réordonne son
#: classeur : elle est donc nommée, et le décompte obtenu est contrôlé.
FEUILLE_TARIF = "ورقة1"
POSITIONS_ATTENDUES_MIN = 5000


def extract_positions(filepath: str) -> List[Dict]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if FEUILLE_TARIF not in wb.sheetnames:
        raise SystemExit(
            f"feuille {FEUILLE_TARIF!r} absente du classeur (présentes : {wb.sheetnames}) — "
            "le tarif a changé de forme, relire la source avant de collecter"
        )
    ws = wb[FEUILLE_TARIF]
    logger.info(f"XLSX: feuille {FEUILLE_TARIF}, {ws.max_row} rows x {ws.max_column} cols")

    positions = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 4:
            continue

        # Colonne 1 = code SH (رمز النظام المنسق)
        hs_code = str(row[1] or "").strip() if row[1] else ""
        if not hs_code or len(hs_code) < 6:
            continue

        code_clean = re.sub(r"[^0-9]", "", hs_code)
        if len(code_clean) < 6:
            continue

        # Colonne 2 = désignation (بيان المنتجات)
        designation = str(row[2] or "").strip() if row[2] else ""

        # Colonne 3 = taux DD (فئة الضريبة)
        dd_raw = row[3]
        dd_rate, dd_interdit = lire_cellule_droit(dd_raw)

        # Colonne 4 = colonne préférentielle « États de la Ligue des États arabes »
        arab_pref_raw = row[4] if len(row) > 4 else None
        arab_pref_rate, arab_interdit = lire_cellule_droit(arab_pref_raw)

        # Construire les taxes
        # Le libellé ne tranche pas entre les dénominations. La loi libyenne en
        # emploie plusieurs pour ce même prélèvement — « الضرائب الجمركية »
        # (28 fois), « الرسوم الجمركية » (15), « الضريبة الجمركية » (4), et la
        # forme cumulative « الضرائب والرسوم الجمركية » (8). Son article des
        # définitions dit que le tarif porte « les taux des DROITS DE DOUANE » ;
        # son article sur la perception dit que ce sont « les TAXES DOUANIÈRES »
        # qui sont perçues. Consacrer l'une des deux serait choisir à la place
        # de la source. Le libellé renvoie donc à la colonne du tarif, dont
        # l'en-tête arabe est reproduit tel quel.
        # Voir LBY_nature_taxe_douaniere_2026-09-18.json.
        # Une cellule VIDE n'est ni une exonération ni une interdiction : c'est
        # une absence. Ne rien émettre ferait servir la position comme si aucun
        # prélèvement n'était dû — un total « complet » amputé de son droit.
        # La ligne est donc émise SANS taux, pour que le moteur la déclare
        # indisponible au lieu de la taire. Relevé sur le tarif 2022 : 9
        # positions dans ce cas, sur 5 920.
        cellule_vide = dd_rate is None and not dd_interdit
        taxes = []
        if cellule_vide:
            taxes.append({
                "code": "DD",
                "name": "Prélèvement du tarif douanier (colonne « فئة الضريبة »)",
                "name_fr": "Prélèvement du tarif douanier — taux absent de la source",
                "name_en": "Customs tariff levy — rate missing from source",
                "name_ar": "فئة الضريبة",
                "rate_pct": None,
                "rate_decimal": None,
                "raw_value": str(dd_raw),
                "base": "CIF",
                "source": "customs.gov.ly",
                "legal_ref": None,
                "is_customs_duty": True,
                "is_vat": False,
                "is_excise": False,
                "note": (
                    "Cellule vide dans le tarif officiel : le taux n'est pas publié "
                    "pour cette position. Ni exonération, ni interdiction — une absence."
                ),
            })
        if dd_rate is not None:
            taxes.append({
                "code": "DD",
                "name": "Prélèvement du tarif douanier (colonne « فئة الضريبة »)",
                "name_fr": "Prélèvement du tarif douanier (colonne « فئة الضريبة »)",
                "name_en": "Customs tariff levy (column « فئة الضريبة »)",
                "name_ar": "فئة الضريبة",
                "denominations_legales": [
                    "الضرائب الجمركية (taxes douanières)",
                    "الرسوم الجمركية (droits de douane)",
                    "الضرائب والرسوم الجمركية (taxes et droits de douane)",
                ],
                "rate_pct": dd_rate,
                "rate_decimal": dd_rate / 100,
                "raw_value": str(dd_raw),
                "base": "CIF",
                "source": "customs.gov.ly",
                "legal_ref": None,
                "is_customs_duty": True,
                "is_vat": False,
                "is_excise": False,
            })

        # Colonne préférentielle. Le tarif la PUBLIE position par position : ce
        # n'est pas une franchise déduite d'une appartenance à un bloc, mais
        # un taux national, au même titre que le droit NPF de la colonne
        # précédente. L'appartenance de l'origine à la Ligue arabe reste, elle, une
        # condition que ce fichier ne tranche pas.
        preferential_rates = []
        if arab_pref_rate is not None:
            preferential_rates.append({
                # L'en-tête de la colonne dit « التعريفة التفضيلية لدول جامعة
                # الدول العربية » — tarif préférentiel pour les ÉTATS DE LA LIGUE
                # DES ÉTATS ARABES. Il ne dit pas GZALE/GAFTA, qui est un accord
                # distinct dont ni le tarif ni la loi ne font mention. Nommer
                # cette colonne GZALE lui prêterait un régime que la source
                # n'invoque pas, et une liste de pays que rien n'établit ici.
                "regime": "LIGUE_ARABE",
                "regime_name_fr": "États de la Ligue des États arabes (colonne du tarif)",
                "rate_pct": arab_pref_rate,
                "raw_value": str(arab_pref_raw),
                "source": "customs.gov.ly",
                "colonne_source": "التعريفة التفضيلية لدول جامعة الدول العربية",
            })

        # La colonne préférentielle est aussi émise parmi les taxes, sous le code
        # que le constructeur du socle reconnaît comme régime préférentiel. Sans
        # cela elle resterait dans `preferential_rates`, que le constructeur ne
        # lit pas, et les 5 890 exonérations publiées par le tarif pour les États
        # de la Ligue arabe seraient perdues. Elle n'entre JAMAIS dans la cascade
        # NPF : le socle la range sous son propre régime.
        if arab_pref_rate is not None:
            taxes.append({
                "code": "LIGUE_ARABE",
                "name": "États de la Ligue des États arabes (colonne du tarif)",
                "name_fr": "États de la Ligue des États arabes (colonne du tarif)",
                "name_ar": "التعريفة التفضيلية لدول جامعة الدول العربية",
                "rate_pct": arab_pref_rate,
                "rate_decimal": arab_pref_rate / 100,
                "raw_value": str(arab_pref_raw),
                "base": "CIF",
                "source": "customs.gov.ly",
                "is_customs_duty": False,
                "is_vat": False,
                "is_excise": False,
                "note": (
                    "Colonne publiée par le tarif de destination. L'appartenance de "
                    "l'origine aux États de la Ligue arabe et les règles d'origine "
                    "applicables ne sont PAS tranchées par ce fichier."
                ),
            })

        # Une interdiction d'importation n'est pas un droit : elle est portée
        # comme restriction de la position, et le prélèvement correspondant
        # reste absent plutôt que fixé à zéro.
        restrictions = []
        if dd_interdit or arab_interdit:
            restrictions.append({
                "type": "IMPORTATION_INTERDITE",
                "portee": "NPF et Ligue arabe" if (dd_interdit and arab_interdit) else (
                    "NPF" if dd_interdit else "Ligue arabe"
                ),
                "verbatim": MENTION_INTERDIT,
                "libelle_fr": "Importation interdite",
                "source": "customs.gov.ly (التعريفة الجمركية 2022)",
            })

        positions.append({
            "national_code": code_clean,
            "hs6": code_clean[:6],
            "chapter": code_clean[:2],
            "heading": code_clean[:4] + "." + code_clean[4:6],
            "section": "",
            "statistical_unit": "",
            "check_digit": "",
            "designation": {
                "fr": "",
                "en": "",
                "ar": designation,
                "full_fr": "",
                "verbatim": hs_code,
            },
            "taxes": taxes,
            "export_taxes": [],
            "preferential_rates": preferential_rates,
            "fiscal_advantages": [],
            "formalities": [],
            "restrictions": restrictions,
            "legal_refs": [],
            "reglementation": {"import": [], "export": []},
            "quotas": {"qcs": None, "qci": None},
            "zlecaf_schedule": {"applied": False, "rate_pct": None, "instruction": None},
            "source_gaps": [],
            "lf_provisions": None,
            "data_status": "crawled_authentic",
            "source_quality": "crawled_authentic",
            "source": "customs.gov.ly (التعريفة الجمركية 2022)",
            "source_url": XLSX_URL,
            "raw_data": {"hs_code": hs_code, "designation_ar": designation, "dd_raw": str(dd_raw), "arab_pref_raw": str(arab_pref_raw)},
        })

    wb.close()
    if len(positions) < POSITIONS_ATTENDUES_MIN:
        raise SystemExit(
            f"{len(positions)} positions collectées, moins que les "
            f"{POSITIONS_ATTENDUES_MIN} attendues — collecte refusée plutôt que "
            "servir un tarif amputé (voir le piège de la feuille ci-dessus)"
        )
    return positions


#: Deux prélèvements que le TARIF NE PORTE PAS, mais que la collecte précédente
#: servait sur chaque position depuis PwC. Ils sont reportés tels quels — mêmes
#: taux, même source, même réserve — pour qu'aucun ne disparaisse du calcul au
#: passage du référentiel WITS au tarif national.
#:
#: Ils restent SANS ASSIETTE, comme avant, et donc non liquidables. Ce n'est pas
#: un oubli : la fiche LBY_assiette_TSP_TP_2026-09-17.json relève elle-même que
#: « TSP 4 % / TP 2 % du crawl DIVERGENT de la service fee (4-5 %) et des
#: "autres droits" (0,5 %) documentés ». Leur donner une assiette les rendrait
#: liquidables et facturerait six points sur une base que notre propre fiche
#: juge douteuse ; les supprimer amputerait le total s'ils sont réels. Ils sont
#: donc portés, visibles, et non liquidés — jusqu'à ce qu'une source tranche.
COMPLEMENTS_NATIONAUX = [
    {
        "code": "TSP",
        "name": "Port Services Tax (taxe des services portuaires)",
        "rate_pct": 4.0,
        "raw_value": "4.0 %",
        "source": "PwC Worldwide Tax Summaries — Libya",
        "source_url": "https://taxsummaries.pwc.com/libya/corporate/other-taxes",
        "as_of": "2026",
        "note": (
            "Taux national standard, non vérifié position par position, ABSENT du "
            "tarif douanier 2022. Assiette non établie : voir la fiche "
            "LBY_assiette_TSP_TP_2026-09-17.json, qui relève une divergence avec la "
            "service fee (4-5 %) documentée par ailleurs."
        ),
    },
    {
        "code": "TP",
        "name": "Production Tax (taxe de production)",
        "rate_pct": 2.0,
        "raw_value": "2.0 %",
        "source": "PwC Worldwide Tax Summaries — Libya",
        "source_url": "https://taxsummaries.pwc.com/libya/corporate/other-taxes",
        "as_of": "2026",
        "note": (
            "Taux national standard, non vérifié position par position, ABSENT du "
            "tarif douanier 2022. Assiette non établie : voir la fiche "
            "LBY_assiette_TSP_TP_2026-09-17.json, qui relève une divergence avec les "
            "« autres droits » (0,5 %) documentés par ailleurs."
        ),
    },
]


def save(positions: List[Dict]):
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    for position in positions:
        for complement in COMPLEMENTS_NATIONAUX:
            position["taxes"].append({
                **complement,
                "name_fr": complement["name"],
                "base": None,
                "is_customs_duty": False,
                "is_vat": False,
                "is_excise": True,
            })

    all_tax_codes = set()
    for p in positions:
        for t in p.get("taxes", []):
            all_tax_codes.add(t.get("code", ""))

    result = {
        "country": "LBY",
        "country_name": "Libya",
        # L'ordre et les bases. Seul le prélèvement du tarif a une assiette
        # établie — la valeur en douane, définie par la loi libyenne sur les
        # douanes comme la valeur transactionnelle augmentée du transport et de
        # l'assurance jusqu'au point d'entrée, soit le CIF. TSP et TP n'en ont
        # pas, à dessein (voir COMPLEMENTS_NATIONAUX).
        # PAS DE TVA : la Libye n'en a pas. Zéro occurrence de « القيمة المضافة »
        # dans la loi, aucune colonne dans le tarif. Ne pas en ajouter une.
        "calculation_rules": {
            "order": ["DD", "TSP", "TP"],
            "bases": {"DD": {"basis": "CIF", "type": "ad_valorem"}},
            "source": (
                "Tarif douanier national 2022 (customs.gov.ly), colonne « فئة الضريبة », "
                "position par position — remplace la moyenne SH6 WITS/UNCTAD-TRAINS. "
                "Assiette établie par la loi libyenne sur les douanes (valeur en douane = "
                "valeur transactionnelle + transport + assurance jusqu'au point d'entrée) : "
                "voir LBY_nature_taxe_douaniere_2026-09-18.json. Pas de TVA en Libye. "
                "TSP et TP sont des compléments nationaux hors tarif, sans assiette établie."
            ),
        },
        "source": "customs.gov.ly (التعريفة الجمركية 2022)",
        "source_url": XLSX_URL,
        "source_quality": "crawled_authentic",
        "extracted_at": datetime.utcnow().strftime("%Y-%m-%d"),
        "stats": {
            "total_positions": len(positions),
            "unique_tax_codes": sorted(all_tax_codes),
            "chapters_covered": len(set(p["chapter"] for p in positions)),
        },
        "sub_positions": positions,
    }

    output = DATA_DIR / "LBY_tariffs.json"
    with open(output, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    logger.info(f"Saved {len(positions)} positions to {output}")


def main():
    logger.info("=== Libya Tariff Scraper ===")
    filepath = download_xlsx()
    if not filepath:
        return

    positions = extract_positions(filepath)
    logger.info(f"Extracted {len(positions)} positions")

    if positions:
        save(positions)

        from collections import Counter
        dd_dist = Counter()
        for p in positions:
            for t in p["taxes"]:
                if t["code"] == "DD":
                    dd_dist[t["rate_pct"]] += 1
        print(f"\nDD distribution: {dict(sorted(dd_dist.items(), key=lambda x: -x[1]))}")
        print(f"Chapters: {len(set(p['chapter'] for p in positions))}")
        print(f"Sample: {positions[0]['national_code']} {positions[0]['designation']['ar'][:40]}")


if __name__ == "__main__":
    main()
